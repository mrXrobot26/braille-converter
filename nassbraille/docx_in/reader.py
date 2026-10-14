"""Read a .docx into the shared Document model: the printed words, text only,
in logical reading order.

* Body order: w:p, w:tbl, w:sdt/w:sdtContent and w:customXml (recursively).
* One printed line = one Block: soft line breaks (w:br, w:cr) split a paragraph
  into several blocks of the same kind.
* Page boundaries: w:br type="page", w:lastRenderedPageBreak, pageBreakBefore
  and section breaks that start a new page. Adjacent boundaries merge, so no
  empty pages are made; footnotes are appended at the end of their page.
  Inside tables a boundary is placed before or after a whole row (Word's own
  marks there are scattered over the cells of a row that runs over a page).
* Revisions are read as accepted (w:ins kept, w:del / w:moveFrom dropped);
  hidden text, field codes, comments, headers and footers are ignored.
* Text boxes become "box" blocks (one box_id per box) placed after their anchor
  paragraph, or after the list that follows an anchor ending with ":" or an
  anchor that is itself a list item.
* Pictures become one "picture" block each; charts and SmartArt are skipped
  with a page warning.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from lxml import etree

from ..model import PICTURE_MARKER, Block, Document, Page
from ..textnorm import clean
from .drawing import Found, scan_drawing, scan_object, scan_pict
from .ns import W, attr, children, int_val, local, on_off, val
from .numbering import NumberingResolver, format_number
from .styles import StyleSheet, heading_kind, on_off_attr

WARN_PICTURES = "فيه صور اتشالت"
WARN_CHART = "فيه رسم بياني اتشال"
WARN_SMARTART = "فيه رسم SmartArt اتشال"
WARN_EQUATION = "فيه معادلة اتشالت"
WARN_SYMBOL = "فيه رموز خاصة (Symbol/Wingdings) اتشالت"
WARN_ALTCHUNK = "فيه جزء متضمّن من ملف تاني ما اتقراش"

MSG_READING = "بقرا ملف الوورد..."
MSG_DONE = "خلصت قراية ملف الوورد"


def _w(name):
    return W + name


W_BODY, W_P, W_TBL, W_TR, W_TC = _w("body"), _w("p"), _w("tbl"), _w("tr"), _w("tc")
W_R, W_T, W_PPR, W_RPR, W_TCPR, W_TRPR = _w("r"), _w("t"), _w("pPr"), _w("rPr"), _w("tcPr"), _w("trPr")
W_SDT, W_SDTCONTENT, W_CUSTOMXML = _w("sdt"), _w("sdtContent"), _w("customXml")
W_SECTPR, W_ALTCHUNK = _w("sectPr"), _w("altChunk")
W_TAB, W_PTAB, W_BR, W_CR = _w("tab"), _w("ptab"), _w("br"), _w("cr")
W_NBHYPHEN, W_SOFTHYPHEN, W_SYM = _w("noBreakHyphen"), _w("softHyphen"), _w("sym")
W_FLDCHAR, W_LRPB = _w("fldChar"), _w("lastRenderedPageBreak")
W_FN_REFERENCE, W_EN_REFERENCE = _w("footnoteReference"), _w("endnoteReference")
W_FN_REF, W_EN_REF = _w("footnoteRef"), _w("endnoteRef")
W_DRAWING, W_PICT, W_OBJECT = _w("drawing"), _w("pict"), _w("object")
W_RUBY, W_RUBYBASE = _w("ruby"), _w("rubyBase")
W_VANISH, W_RSTYLE, W_RFONTS, W_B = _w("vanish"), _w("rStyle"), _w("rFonts"), _w("b")
W_DEL, W_INS, W_FRAMEPR = _w("del"), _w("ins"), _w("framePr")
W_FOOTNOTE, W_ENDNOTE = _w("footnote"), _w("endnote")
M_NS = "{%s}" % "http://schemas.openxmlformats.org/officeDocument/2006/math"
M_OMATH, M_OMATHPARA, M_T = M_NS + "oMath", M_NS + "oMathPara", M_NS + "t"

INLINE_RECURSE = {_w(n) for n in ("ins", "moveTo", "hyperlink", "smartTag", "customXml", "dir",
                                  "bdo", "fldSimple")}
INLINE_SKIP = {_w(n) for n in ("pPr", "rPr", "del", "moveFrom", "sdtPr", "sdtEndPr", "subDoc",
                               "commentRangeStart", "commentRangeEnd", "bookmarkStart",
                               "bookmarkEnd", "proofErr", "permStart", "permEnd")}
MATH_PLAIN = {"oMathPara", "oMathParaPr", "oMath", "r", "t", "rPr", "ctrlPr", "jc", "sty", "scr",
              "nor", "lit", "brk", "aln", "argPr", "argSz"}
TEXT_KINDS = {"h1", "h2", "h3", "h4", "para", "box", "caption", "notice"}

# Manually typed list items: "- ", "• ", "1- ", "1.", "(1)", "أ- ", "(أ)" ...
MANUAL_LIST = re.compile(
    r"^(?:[-–—•▪◦●○■□◆◇❖➢►✓✔*]\s*\S"
    r"|\(?[0-9٠-٩]{1,3}\s*[-–.)/]"
    r"|\(?[ء-ي]\s*[-–)]"
    r"|\([0-9٠-٩ء-ي]{1,3}\))"
)

# Symbol-font glyphs used as bullets (code points without the 0xF000 offset).
_BULLETS = {
    "symbol": {0xB7, 0xA8, 0xAE, 0xDE, 0xE0, 0x2D, 0x2A},
    "wingdings": {0x6C, 0x6E, 0x6F, 0x70, 0x71, 0x72, 0x73, 0x74, 0x75, 0x76, 0x77, 0x9F,
                  0xA1, 0xA7, 0xA8, 0xD8, 0xDF, 0xE0, 0xE8, 0xF0, 0xFC, 0xFE, 0x46},
    "wingdings 2": {0x50, 0x52, 0x97, 0x98, 0x99, 0x9A, 0x9E, 0xA2, 0xA3, 0xA4, 0xB7, 0xB8},
    "wingdings 3": {0x22, 0x67, 0x75, 0x7D, 0x84, 0xC6, 0xD8},
    "webdings": {0x34, 0x38, 0x3C, 0x3D, 0x69, 0x6E},
}
UNICODE_BULLETS = set("•●○◦▪▫■□◆◇❖➢►▶✓✔✗‣⁃∙⦁")
PUA = re.compile("[\uf000-\uf0ff]")
# raw line breaks / tabs typed inside w:t are shown by Word as spaces
_CONTROL_TO_SPACE = str.maketrans({"\n": " ", "\r": " ", "\t": " ", "\v": " ", "\f": " "})

_PAGE = object()          # page boundary marker inside collected item lists
LINE, PAGE, PIC = ("line",), ("page",), ("pic",)
SPACE = ("t", " ")


# -- small sinks ---------------------------------------------------------------
class _ListSink:
    """Collects blocks and page boundaries (table cells, text boxes, notes).

    A table cell (in_cell) only remembers whether a page boundary came before
    its text or inside it: tables place boundaries per row, so a cell's line
    is not cut where Word's page happened to break it.
    """

    def __init__(self, in_cell=False):
        self.items = []
        self.in_cell = in_cell
        self.page_start = self.page_inside = False

    def block(self, block):
        self.items.append(block)

    def page(self, at_start=None):
        if not self.in_cell:
            self.items.append(_PAGE)
            return
        if at_start is None:
            at_start = not self.items
        if at_start:
            self.page_start = True
        else:
            self.page_inside = True

    def blocks(self):
        return [b for b in self.items if b is not _PAGE]


class _Pages:
    """The page builder: merges adjacent boundaries, appends footnotes."""

    def __init__(self, on_new_page=None):
        self.done = []
        self.blocks, self.warnings, self.notes = [], [], []
        self.on_new_page = on_new_page

    def block(self, block):
        self.blocks.append(block)

    def warn(self, message):
        if message not in self.warnings:
            self.warnings.append(message)

    def note(self, blocks):
        self.notes.extend(blocks)

    def page(self):
        if self.blocks or self.warnings or self.notes:
            self._close()
            if self.on_new_page is not None:
                self.on_new_page()

    def _close(self):
        self.done.append(Page(number=len(self.done) + 1, blocks=self.blocks + self.notes,
                              warnings=self.warnings))
        self.blocks, self.warnings, self.notes = [], [], []

    def finish(self, tail=()):
        tail = list(tail)
        if self.blocks or self.warnings or self.notes or not self.done:
            self.notes.extend(tail)
            if self.blocks or self.warnings or self.notes:
                self._close()
        elif tail:
            self.done[-1].blocks.extend(tail)
        return self.done


class _Story:
    """Parse state of one story: the body, a text box, a footnote."""

    __slots__ = ("mode", "box_id", "fields", "broken", "note_mark", "dropcap")

    def __init__(self, mode, container, box_id=None, note_mark=None):
        self.mode = mode            # body | box | note | art
        self.box_id = box_id
        self.fields = []            # complex-field phases: "instr" | "result"
        self.broken = _unclosed_fields(container)
        self.note_mark = note_mark  # printed number of the note being read
        self.dropcap = None


_SCAN_SKIP = {W + "del", W + "moveFrom", W + "txbxContent"}


def _unclosed_fields(container):
    """fldChar "begin" marks of one story that never reach their "end".

    They are ignored, so a damaged field cannot hide the rest of the text.
    The walk mirrors the reader: mc:Choice only, no deleted text, text boxes
    are stories of their own.
    """
    stack = []

    def walk(el):
        for ch in children(el):
            tag = ch.tag
            if tag in _SCAN_SKIP:
                continue
            if tag == W + "fldChar":
                kind = ch.get(W + "fldCharType")
                if kind == "begin":
                    stack.append(ch)
                elif kind == "end" and stack:
                    stack.pop()
            else:
                walk(ch)

    if container is not None:
        walk(container)
    return set(stack)


class _Para:
    __slots__ = ("events", "heading", "label", "page_before", "page_after", "section_end",
                 "has_text", "has_objects", "ends_colon", "is_list", "dropcap")


class _Cell:
    __slots__ = ("start", "span", "cont", "tc", "blocks", "lines", "extras", "text", "words",
                 "page_start", "page_inside")

    def __init__(self, start, span, cont, tc):
        self.start, self.span, self.cont, self.tc = start, span, cont, tc
        self.blocks, self.lines, self.extras = [], [], []
        self.text, self.words = "", 0
        self.page_start = self.page_inside = False   # page boundary before / within the cell


class _Row:
    __slots__ = ("cells", "header")

    def __init__(self, cells, header):
        self.cells, self.header = cells, header


# -- public entry point ---------------------------------------------------------
def read_docx(path, progress=None):
    """Read a .docx file; progress(done, total, message) is optional."""
    path = Path(path)
    try:
        package, part = _open_package(path)
    except Exception as exc:
        if not _legacy_signature(path):
            raise
        # an old .doc / RTF file that only carries a .docx name: Word opens those
        # without complaint, so convert it like any .doc
        from ..doc_legacy import convert_to_docx
        try:
            package, part = _open_package(convert_to_docx(path))
        except Exception:
            raise exc from None
    reader = _Reader(package, part, progress)
    pages = reader.run()
    return Document(source=str(path), kind="docx", pages=pages, report=reader.report(pages))


def _conversion_error(message):
    from ..convert import ConversionError
    return ConversionError(message)


def _open_package(path):
    import docx  # noqa: F401  (registers python-docx's part classes)
    from docx.package import Package

    try:
        package = Package.open(str(path))
    except Exception as exc:
        raise _conversion_error(_bad_file_message(path)) from exc
    try:
        part = package.main_document_part
    except Exception as exc:
        if _is_strict(path):
            raise _conversion_error(
                "الملف ده متسجّل بصيغة Strict Open XML. افتحه في Word واحفظه تاني "
                "بصيغة \"مستند Word\" العادية (docx).") from exc
        raise _conversion_error(_bad_file_message(path)) from exc
    return package, part


def _bad_file_message(path):
    try:
        with open(path, "rb") as fh:
            head = fh.read(8)
    except PermissionError:
        return "الملف مقفول أو مفتوح في برنامج تاني. اقفله في Word وجرّب تاني."
    except OSError:
        head = b""
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return ("الملف ده عليه كلمة سر أو متسجّل بصيغة Word قديمة. افتحه في Word، "
                "شيل كلمة السر لو موجودة، واحفظه بصيغة docx.")
    return "الملف ده مش ملف Word سليم. افتحه في Word واحفظه تاني بصيغة docx."


def _legacy_signature(path):
    """Word 97-2003 (OLE) or RTF content, whatever the file name says."""
    try:
        with open(path, "rb") as fh:
            head = fh.read(8)
    except OSError:
        return False
    return head.startswith(b"\xd0\xcf\x11\xe0") or head.startswith(b"{\\rtf")


def _is_strict(path):
    try:
        with zipfile.ZipFile(path) as zf:
            return b"purl.oclc.org/ooxml" in zf.read("_rels/.rels")
    except (OSError, KeyError, zipfile.BadZipFile):
        return False


_PARSER = etree.XMLParser(resolve_entities=False, huge_tree=True)


def _xml(part):
    element = getattr(part, "element", None)
    if element is not None:
        return element
    return etree.fromstring(part.blob, _PARSER)


def _related(part):
    out = {}
    for rel in part.rels.values():
        if rel.is_external:
            continue
        name = rel.reltype.rsplit("/", 1)[-1]
        if name in ("styles", "numbering", "footnotes", "endnotes", "settings") and name not in out:
            try:
                out[name] = _xml(rel.target_part)
            except Exception:       # a broken side part should not stop the body text
                continue
    return out


def _notes(root, tag):
    notes = {}
    if root is None:
        return notes
    for el in root.iter(tag):
        if attr(el, "type", "normal") == "normal":
            notes.setdefault(attr(el, "id"), el)
    return notes


def _sym_font(name):
    """The legacy (non-Unicode) symbol font family of a font name, or None."""
    name = (name or "").strip().lower()
    for family in sorted(_BULLETS, key=len, reverse=True):
        if name.startswith(family):
            return family
    return None


def _bullet_glyph(family, code):
    return code in _BULLETS.get(family, ())


# -- the reader -------------------------------------------------------------------
class _Reader:
    def __init__(self, package, part, progress=None):
        self.progress = progress
        self.root = _xml(part)
        self.body = self.root.find(W_BODY)
        related = _related(part)
        self.styles = StyleSheet(related.get("styles"))
        self.numbering = NumberingResolver(related.get("numbering"), self.styles)
        self.footnotes = _notes(related.get("footnotes"), W_FOOTNOTE)
        self.endnotes = _notes(related.get("endnotes"), W_ENDNOTE)
        final_sect = self.body.find(W_SECTPR) if self.body is not None else None
        self.fn_fmt, self.fn_start, self.fn_restart = _note_numbering(
            related.get("settings"), final_sect, "footnotePr", "decimal")
        self.en_fmt, self.en_start, _ = _note_numbering(
            related.get("settings"), final_sect, "endnotePr", "lowerRoman")
        self.fn_next, self.en_next = self.fn_start, self.en_start
        self.pages = _Pages(on_new_page=self._new_page)
        self.endnote_blocks = []
        self.box_seq = 0
        self.obj_seq = 0
        self.stats = {"pictures": 0, "text_boxes": 0, "tables": 0, "footnotes": 0,
                      "endnotes": 0, "charts": 0, "smartart": 0, "equations": 0,
                      "symbols_dropped": 0, "hidden_runs": 0}
        self.sect_breaks = self._section_breaks()
        self.total = 1

    def run(self):
        if self.body is not None:
            self._walk(self.body, _Story("body", self.body), self.pages, top=True)
        pages = self.pages.finish(self.endnote_blocks)
        if self.progress is not None:
            self.progress(self.total, self.total, MSG_DONE)
        return pages

    def report(self, pages):
        out = {"reader": "docx", "pages": len(pages),
               "blocks": sum(len(p.blocks) for p in pages)}
        out.update(self.stats)
        return out

    def _new_page(self):
        if self.fn_restart == "eachPage":
            self.fn_next = self.fn_start

    # -- containers ------------------------------------------------------------
    def _iter_block(self, container):
        for ch in children(container):
            tag = ch.tag
            if tag in (W_P, W_TBL, W_ALTCHUNK):
                yield ch
            elif tag == W_SDT:
                content = ch.find(W_SDTCONTENT)
                if content is not None:
                    yield from self._iter_block(content)
            elif tag == W_CUSTOMXML:
                yield from self._iter_block(ch)

    def _section_breaks(self):
        """{paragraph-level sectPr element: does a new page follow it?}

        The next section's w:type decides (default nextPage); a "continuous"
        section that changes the page size or orientation still starts a page.
        """
        if self.body is None:
            return {}
        sects = []
        for el in self._iter_block(self.body):
            if el.tag != W_P:
                continue
            ppr = el.find(W_PPR)
            sect = ppr.find(W_SECTPR) if ppr is not None else None
            if sect is not None:
                sects.append(sect)
        sects.append(self.body.find(W_SECTPR))
        breaks = {}          # the dict keeps the element proxies alive, so lookups match
        for prev, nxt in zip(sects, sects[1:]):
            kind = val(nxt.find(_w("type")), "nextPage") if nxt is not None else "nextPage"
            before, after = _page_size(prev), _page_size(nxt)
            resized = before is not None and after is not None and before != after
            breaks[prev] = kind not in ("continuous", "nextColumn") or resized
        return breaks

    def _walk(self, container, story, sink, top=False):
        pending = []          # box blocks waiting for the end of a list
        elements = list(self._iter_block(container))
        total = len(elements)
        step = max(1, total // 50)
        if top:
            self.total = max(1, total)
        for i, el in enumerate(elements):
            if top and self.progress is not None and i % step == 0:
                self.progress(i, total, MSG_READING)
            tag = el.tag
            if tag == W_P:
                para = self._read_para(el, story)
                if para.dropcap is not None:
                    story.dropcap = (story.dropcap or "") + para.dropcap
                    continue
                blank = not (para.has_text or para.label or para.has_objects)
                if pending and not (para.is_list or blank):
                    self._flush(pending, sink)
                if para.page_before:
                    self._flush(pending, sink)
                    sink.page()
                self._emit_para(para, story, sink, pending)
                if para.page_after:
                    self._flush(pending, sink)
                    sink.page()
                if para.section_end and self.fn_restart == "eachSect":
                    self.fn_next = self.fn_start
            elif tag == W_TBL:
                self._flush(pending, sink)
                self._table(el, story, sink)
            elif tag == W_ALTCHUNK:
                self.pages.warn(WARN_ALTCHUNK)
        self._flush(pending, sink)
        if story.dropcap:                       # a drop cap with nothing after it
            sink.block(self._line_block(story, None, story.dropcap))
            story.dropcap = None

    @staticmethod
    def _flush(pending, sink):
        for block in pending:
            sink.block(block)
        pending.clear()

    # -- paragraphs --------------------------------------------------------------
    def _read_para(self, p, story):
        ppr = p.find(W_PPR)
        pstyle = self.styles.para_style_id(p)
        para = _Para()
        para.heading = None
        if story.mode == "body":
            para.heading = heading_kind(self.styles.outline_level(ppr, pstyle))
        mark_rpr = ppr.find(W_RPR) if ppr is not None else None
        mark_gone = mark_rpr is not None and (
            mark_rpr.find(W_DEL) is not None or bool(on_off(mark_rpr.find(W_VANISH))))
        para.label = "" if mark_gone else self.numbering.label_for(p)
        para.page_before = story.mode == "body" and self._page_break_before(ppr, pstyle)
        para.page_after = para.section_end = False
        sect = ppr.find(W_SECTPR) if story.mode == "body" and ppr is not None else None
        if sect is not None and sect in self.sect_breaks:
            para.section_end = True
            para.page_after = self.sect_breaks[sect]
        events = []
        self._inline(p, story, events, pstyle)
        para.events = events
        lines = _line_texts(events)
        texts = [clean(t) for t in lines]
        texts = [t for t in texts if t]
        para.has_text = bool(texts)
        if mark_gone and para.has_text:
            # deleted/hidden paragraph mark but visible text: it still prints, numbered
            para.label = self.numbering.label_for(p)
        para.has_objects = any(e[0] in ("obj", "pic") for e in events)
        para.ends_colon = bool(texts) and texts[-1].endswith(":")
        first = texts[0] if texts else ""
        para.is_list = bool(para.label.strip()) or bool(MANUAL_LIST.match(first))
        para.dropcap = None
        frame = ppr.find(W_FRAMEPR) if ppr is not None else None
        if frame is not None and attr(frame, "dropCap") in ("drop", "margin") and para.has_text:
            para.dropcap = "".join(lines).strip()
        return para

    def _page_break_before(self, ppr, pstyle):
        el = ppr.find(_w("pageBreakBefore")) if ppr is not None else None
        if el is None:
            el = self.styles.ppr_child(pstyle, _w("pageBreakBefore"))
        return bool(on_off(el))

    def _emit_para(self, para, story, sink, pending):
        cur = []
        if story.dropcap:
            cur.append(story.dropcap)
            story.dropcap = None
        label = [para.label]
        late_pics = [0]
        objs = []

        def flush(force_label):
            text = "".join(cur)
            cur.clear()
            if label[0] and (force_label or clean(text)):
                text = label[0] + text
                label[0] = ""
            if clean(text):
                sink.block(self._line_block(story, para.heading, text))
            for _ in range(late_pics[0]):
                sink.block(self._picture(story))
            late_pics[0] = 0

        for ev in para.events:
            kind = ev[0]
            if kind == "t":
                cur.append(ev[1])
            elif kind == "line":
                flush(True)
            elif kind == "page":
                if getattr(sink, "in_cell", False):     # tables break pages per row
                    sink.page(at_start=not sink.items and not clean("".join(cur)))
                    continue
                flush(False)
                for block in self._ordered(objs):
                    sink.block(block)
                objs.clear()
                self._flush(pending, sink)
                sink.page()
            elif kind == "pic":
                if not clean("".join(cur)) and not label[0]:
                    sink.block(self._picture(story))
                else:
                    late_pics[0] += 1
            elif kind == "obj":
                objs.append(ev)
            elif kind == "fn":
                self._attach_note(ev[1], ev[2], ev[3])
            elif kind == "warn":
                self.pages.warn(ev[1])
        flush(True)
        if objs:
            ordered = self._ordered(objs)
            if para.ends_colon or para.is_list:
                pending.extend(ordered)        # keep a ":" list together
            else:
                for block in ordered:
                    sink.block(block)

    @staticmethod
    def _ordered(objs):
        out = []
        for ev in sorted(objs, key=lambda e: e[1]):
            out.extend(ev[2])
        return out

    @staticmethod
    def _kind(story, heading):
        if story.mode == "box":
            return "box"
        if story.mode == "body" and heading:
            return heading
        return "para"

    def _line_block(self, story, heading, text):
        kind = self._kind(story, heading)
        return Block(kind, text.strip(), box_id=story.box_id if kind == "box" else None)

    @staticmethod
    def _picture(story):
        return Block("picture", PICTURE_MARKER, box_id=story.box_id)

    # -- inline content ------------------------------------------------------------
    def _inline(self, el, story, ev, pstyle):
        for ch in children(el):
            tag = ch.tag
            if tag == W_R:
                self._run(ch, story, ev, pstyle)
            elif tag in INLINE_RECURSE:
                self._inline(ch, story, ev, pstyle)
            elif tag == W_SDT:
                content = ch.find(W_SDTCONTENT)
                if content is not None:
                    self._inline(content, story, ev, pstyle)
            elif tag in (M_OMATH, M_OMATHPARA):
                if "instr" not in story.fields:
                    self._math(ch, ev)
            elif tag in INLINE_SKIP:
                continue
            elif tag.startswith(W) and not tag.endswith("Pr") and not tag.endswith("PrChange"):
                self._inline(ch, story, ev, pstyle)   # unknown wrapper: look inside

    def _run(self, r, story, ev, pstyle):
        rpr = r.find(W_RPR)
        hidden = self._hidden(rpr, pstyle)
        if hidden:
            self.stats["hidden_runs"] += 1
        for ch in children(r):
            tag = ch.tag
            if tag == W_FLDCHAR:
                kind = attr(ch, "fldCharType")
                if kind == "begin":
                    if ch not in story.broken:
                        story.fields.append("instr")
                elif kind == "separate" and story.fields:
                    story.fields[-1] = "result"
                elif kind == "end" and story.fields:
                    story.fields.pop()
                continue
            if tag == W_LRPB:
                if story.mode == "body":
                    ev.append(PAGE)
                continue
            if hidden or "instr" in story.fields:
                continue
            if tag == W_T:
                text = (ch.text or "").translate(_CONTROL_TO_SPACE)
                if text and PUA.search(text):
                    text = self._symbol_text(text, rpr, pstyle, ev)
                if text:
                    ev.append(("t", text))
            elif tag in (W_TAB, W_PTAB):
                ev.append(SPACE)
            elif tag == W_BR:
                if attr(ch, "type") == "page" and story.mode == "body":
                    ev.append(PAGE)
                else:
                    ev.append(LINE)
            elif tag == W_CR:
                ev.append(LINE)
            elif tag == W_NBHYPHEN:
                ev.append(("t", "-"))
            elif tag == W_SYM:
                text = self._sym(ch)
                if text:
                    ev.append(("t", text))
                else:
                    self.stats["symbols_dropped"] += 1
                    ev.append(("warn", WARN_SYMBOL))
            elif tag in (W_FN_REFERENCE, W_EN_REFERENCE):
                self._note_ref(ch, "footnote" if tag == W_FN_REFERENCE else "endnote", story, ev)
            elif tag in (W_FN_REF, W_EN_REF):
                if story.note_mark:
                    ev.append(("t", story.note_mark))
            elif tag == W_DRAWING:
                found = Found()
                scan_drawing(ch, found)
                self._found(found, story, ev)
            elif tag == W_PICT:
                found = Found()
                scan_pict(ch, found)
                self._found(found, story, ev)
            elif tag == W_OBJECT:
                found = Found()
                scan_object(ch, found)
                self._found(found, story, ev)
            elif tag == W_RUBY:
                base = ch.find(W_RUBYBASE)
                if base is not None:
                    self._inline(base, story, ev, pstyle)
            # w:softHyphen, w:instrText, w:delText, separators, comment marks,
            # date/page blocks: nothing printed in the body

    def _hidden(self, rpr, pstyle):
        if rpr is not None:
            el = rpr.find(W_VANISH)
            if el is not None:
                return bool(on_off(el))
            rstyle = val(rpr.find(W_RSTYLE))
            if rstyle:
                el = self.styles.rpr_child(rstyle, W_VANISH, defaults=False)
                if el is not None:
                    return bool(on_off(el))
        return bool(on_off(self.styles.rpr_child(pstyle, W_VANISH)))

    def _run_font(self, rpr, pstyle):
        el = rpr.find(W_RFONTS) if rpr is not None else None
        if el is None and rpr is not None:
            rstyle = val(rpr.find(W_RSTYLE))
            if rstyle:
                el = self.styles.rpr_child(rstyle, W_RFONTS, defaults=False)
        if el is None:
            el = self.styles.rpr_child(pstyle, W_RFONTS)
        if el is None:
            return None
        for key in ("ascii", "hAnsi", "cs", "eastAsia"):
            family = _sym_font(attr(el, key))
            if family:
                return family
        return None

    def _symbol_text(self, text, rpr, pstyle, ev):
        """Private-use characters typed in a symbol font: bullets -> "-", others dropped."""
        family = self._run_font(rpr, pstyle)
        if family is None:
            return text

        def repl(m):
            code = ord(m.group(0)) - 0xF000
            if _bullet_glyph(family, code):
                return "-"
            self.stats["symbols_dropped"] += 1
            ev.append(("warn", WARN_SYMBOL))
            return ""

        return PUA.sub(repl, text)

    def _sym(self, el):
        """w:sym: bullet glyphs -> "-"; other symbol-font glyphs are dropped.

        In a Unicode font (e.g. «ﷺ» inserted from Traditional Arabic) the code is
        a real character and is kept.
        """
        family = _sym_font(attr(el, "font"))
        try:
            code = int(attr(el, "char") or "", 16)
        except ValueError:
            return None
        if family is not None:
            if code >= 0xF000:
                code -= 0xF000
            return "-" if _bullet_glyph(family, code) else None
        if 0xE000 <= code <= 0xF8FF or code > 0x10FFFF:   # private use: unknown glyph
            return None
        ch = chr(code)
        return "-" if ch in UNICODE_BULLETS else ch

    def _math(self, el, ev):
        texts, structured = [], False
        for node in el.iter():
            tag = node.tag
            if not isinstance(tag, str):
                continue
            if tag == M_T:
                texts.append(node.text or "")
            elif tag.startswith(M_NS) and local(tag) not in MATH_PLAIN:
                structured = True
        if structured:
            self.stats["equations"] += 1
            ev.append(("warn", WARN_EQUATION))
            return
        text = "".join(texts)
        if text:
            ev.append(("t", text))

    # -- drawings ----------------------------------------------------------------------
    def _found(self, found, story, ev):
        for (y, x), content, art in found.boxes:
            self.stats["text_boxes"] += 1
            blocks = self._read_box(content, art)
            if blocks:
                self.obj_seq += 1
                ev.append(("obj", (y, -x, self.obj_seq), blocks))
        for (y, x), text in found.wordart:
            blocks = [self._line_block(story, None, line) for line in text.splitlines()
                      if clean(line)]
            if blocks:
                self.obj_seq += 1
                ev.append(("obj", (y, -x, self.obj_seq), blocks))
        if found.pictures:
            self.stats["pictures"] += found.pictures
            ev.append(PIC)
            ev.append(("warn", WARN_PICTURES))
        if found.charts:
            self.stats["charts"] += found.charts
            ev.append(("warn", WARN_CHART))
        if found.smartart:
            self.stats["smartart"] += found.smartart
            ev.append(("warn", WARN_SMARTART))
        if found.equations:
            self.stats["equations"] += found.equations
            ev.append(("warn", WARN_EQUATION))

    def _read_box(self, content, art):
        if art:                                    # modern WordArt: plain text lines
            story = _Story("art", content)
        else:
            self.box_seq += 1
            story = _Story("box", content, box_id=self.box_seq)
        sink = _ListSink()
        self._walk(content, story, sink)
        return sink.blocks()

    # -- notes -------------------------------------------------------------------------
    def _note_ref(self, el, kind, story, ev):
        if story.mode == "note":
            return
        custom = on_off_attr(attr(el, "customMarkFollows"))
        mark = ""
        if not custom:
            if kind == "footnote":
                mark = format_number(self.fn_next, self.fn_fmt)
                self.fn_next += 1
            else:
                mark = format_number(self.en_next, self.en_fmt)
                self.en_next += 1
        if mark:
            ev.append(("t", mark))
        ev.append(("fn", kind, attr(el, "id"), mark))

    def _attach_note(self, kind, note_id, mark):
        source = self.footnotes if kind == "footnote" else self.endnotes
        el = source.get(note_id)
        if el is None:
            return
        story = _Story("note", el, note_mark=mark)
        sink = _ListSink()
        self._walk(el, story, sink)
        blocks = sink.blocks()
        texts = [b for b in blocks if b.kind != "picture"]
        if not texts:
            return
        if mark and not clean(texts[0].text).startswith(mark):
            texts[0].text = mark + " " + texts[0].text
        if kind == "footnote":
            self.stats["footnotes"] += 1
            self.pages.note(blocks)
        else:
            self.stats["endnotes"] += 1
            self.endnote_blocks.extend(blocks)

    # -- tables ------------------------------------------------------------------------
    def _iter_rows(self, tbl):
        for ch in children(tbl):
            if ch.tag == W_TR:
                yield ch
            elif ch.tag == W_SDT:
                content = ch.find(W_SDTCONTENT)
                if content is not None:
                    yield from self._iter_rows(content)
            elif ch.tag == W_CUSTOMXML:
                yield from self._iter_rows(ch)

    def _iter_cells(self, tr):
        for ch in children(tr):
            if ch.tag == W_TC:
                yield ch
            elif ch.tag == W_SDT:
                content = ch.find(W_SDTCONTENT)
                if content is not None:
                    yield from self._iter_cells(content)
            elif ch.tag == W_CUSTOMXML:
                yield from self._iter_cells(ch)

    def _table(self, tbl, story, sink):
        self.stats["tables"] += 1
        rows = []
        for tr in self._iter_rows(tbl):
            trpr = tr.find(W_TRPR)
            if trpr is not None and (trpr.find(W_DEL) is not None
                                     or on_off(trpr.find(_w("hidden")))):
                continue
            header = trpr is not None and bool(on_off(trpr.find(_w("tblHeader"))))
            col = int_val(trpr.find(_w("gridBefore")), 0) if trpr is not None else 0
            cells = []
            for tc in self._iter_cells(tr):
                tcpr = tc.find(W_TCPR)
                span, cont = 1, False
                if tcpr is not None:
                    span = max(1, int_val(tcpr.find(_w("gridSpan")), 1))
                    vmerge, hmerge = tcpr.find(_w("vMerge")), tcpr.find(_w("hMerge"))
                    cont = (vmerge is not None and val(vmerge, "continue") == "continue") or \
                           (hmerge is not None and val(hmerge, "continue") == "continue")
                cell = _Cell(col, span, cont, tc)
                if not cont:
                    collected = _ListSink(in_cell=True)
                    self._walk(tc, story, collected)
                    self._summarize(cell, collected, story)
                col += span
                cells.append(cell)
            if cells:
                rows.append(_Row(cells, header))
        if rows:
            self._format_table(rows, story, sink)

    @staticmethod
    def _summarize(cell, collected, story):
        cell.blocks = collected.blocks()
        cell.page_start, cell.page_inside = collected.page_start, collected.page_inside
        for b in cell.blocks:
            own_box = b.kind != "box" or b.box_id == story.box_id
            if b.kind in TEXT_KINDS and own_box:
                cell.lines.append(b)
            else:
                cell.extras.append(b)
        cell.text = " ".join(t for t in (clean(b.text) for b in cell.lines) if t)
        cell.words = len(cell.text.split())

    @staticmethod
    def _row_breaks(cells):
        """(page before the row, page after the row).

        Word's lastRenderedPageBreak marks in a row that runs over onto the next
        page sit at the split point of some cells and at the start of others, so
        only the first non-empty cell tells whether the row itself starts a new
        page; any other mark means the row runs over: one boundary after it.
        """
        first = next((c for c in cells if c.blocks or c.page_start or c.page_inside), None)
        before = first is not None and first.page_start
        after = any(c.page_inside for c in cells) or \
            (not before and any(c.page_start for c in cells))
        return before, after

    def _format_table(self, rows, story, sink):
        """Rows in order, with at most one page boundary before and after a row."""
        def line(text):
            sink.block(self._line_block(story, None, text))

        def pair(key, cell):
            values = cell.lines if cell is not None and not cell.cont else []
            if key and values:
                sep = " " if key.endswith(":") else ": "
                line(key + sep + values[0].text.strip())
                for b in values[1:]:
                    line(b.text)
            elif key:
                line(key)
            else:
                for b in values:
                    line(b.text)

        def extras(cells):
            for c in cells:
                for b in c.extras:
                    sink.block(b)

        def page_before(cells):
            if self._row_breaks(cells)[0]:
                sink.page()

        def page_after(cells):
            if self._row_breaks(cells)[1]:
                sink.page()

        def key_of(cell):
            return "" if cell.cont else cell.text

        # 1. two rows whose first row is short: one line per column, "row0: row1"
        if len(rows) == 2:
            r0, r1 = rows
            same_grid = [c.start for c in r0.cells] == [c.start for c in r1.cells]
            short = all(c.words <= 4 for c in r0.cells if not c.cont)
            if same_grid and short and not self._key_column(rows):
                before0, after0 = self._row_breaks(r0.cells)
                before1, after1 = self._row_breaks(r1.cells)
                if before0:
                    sink.page()
                for a, b in zip(r0.cells, r1.cells):
                    pair(key_of(a), b)
                extras(r0.cells + r1.cells)
                if after0 or before1 or after1:
                    sink.page()
                return
        # 2. two columns whose first column is short: "c0: c1" per row
        widths = [len(r.cells) for r in rows]
        if max(widths) == 2 and all(r.cells[0].words <= 6 for r in rows
                                    if len(r.cells) == 2 and not r.cells[0].cont):
            for r in rows:
                page_before(r.cells)
                if len(r.cells) == 2:
                    pair(key_of(r.cells[0]), r.cells[1])
                else:
                    for c in r.cells:
                        pair("", c)
                extras(r.cells)
                page_after(r.cells)
            return
        # 3. header-styled first row(s): "header: cell" for every data cell
        n_head = 0
        while n_head < len(rows) and rows[n_head].header:
            n_head += 1
        if 0 < n_head < len(rows):
            heads = {}
            for r in rows[:n_head]:
                page_before(r.cells)
                for c in r.cells:
                    if c.cont or not c.text:
                        continue
                    for col in range(c.start, c.start + c.span):
                        heads[col] = (heads.get(col, "") + " " + c.text).strip()
                extras(r.cells)
            for r in rows[n_head:]:
                page_before(r.cells)
                for c in r.cells:
                    if c.cont or not c.lines:
                        continue
                    pair(heads.get(c.start, ""), c)
                extras(r.cells)
                page_after(r.cells)
            return
        # 4. otherwise: every cell paragraph as it is, rows and cells in order
        for r in rows:
            page_before(r.cells)
            for c in r.cells:
                if not c.cont:
                    for block in c.blocks:
                        sink.block(block)
            page_after(r.cells)

    def _key_column(self, rows):
        """2x2 tables: a bold/shaded first column (and plain first row) reads as
        "key: value" per row rather than per column."""
        if len(rows) != 2 or any(len(r.cells) != 2 for r in rows):
            return False
        c00, c01 = rows[0].cells
        c10 = rows[1].cells[0]
        return (self._emphasized(c00) and self._emphasized(c10)
                and not self._emphasized(c01))

    def _emphasized(self, cell):
        tcpr = cell.tc.find(W_TCPR)
        shd = tcpr.find(_w("shd")) if tcpr is not None else None
        fill = (attr(shd, "fill") or "").lower()
        if fill and fill not in ("auto", "ffffff", "none"):
            return True
        seen = False
        for p in cell.tc.iter(W_P):
            pstyle = self.styles.para_style_id(p)
            for r in p.iter(W_R):
                if not "".join(t.text or "" for t in r.iter(W_T)).strip():
                    continue
                seen = True
                rpr = r.find(W_RPR)
                bold = rpr.find(W_B) if rpr is not None else None
                if bold is None:
                    bold = self.styles.rpr_child(pstyle, W_B)
                if not on_off(bold):
                    return False
        return seen


def _line_texts(events):
    lines, cur = [], []
    for e in events:
        if e[0] == "t":
            cur.append(e[1])
        elif e is LINE or e is PAGE:
            lines.append("".join(cur))
            cur = []
    lines.append("".join(cur))
    return lines


def _page_size(sect):
    size = sect.find(_w("pgSz")) if sect is not None else None
    if size is None:
        return None
    return (attr(size, "w"), attr(size, "h"), attr(size, "orient", "portrait"))


def _note_numbering(settings, sect, tag, default_fmt):
    fmt, start, restart = default_fmt, 1, "continuous"
    for holder in (settings, sect):
        pr = holder.find(_w(tag)) if holder is not None else None
        if pr is None:
            continue
        fmt = val(pr.find(_w("numFmt")), fmt)
        start = int_val(pr.find(_w("numStart")), start)
        restart = val(pr.find(_w("numRestart")), restart)
    return fmt, start, restart


__all__ = ["read_docx", "WARN_PICTURES", "WARN_CHART", "WARN_SMARTART", "WARN_EQUATION",
           "WARN_SYMBOL"]
