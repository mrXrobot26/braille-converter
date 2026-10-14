"""Write a Document as a right-to-left, text-only Word file.

Ported from the script that produced the verified Saifi copy: real RTL
paragraphs (w:bidi), RTL runs, complex-script fonts and sizes, ar-EG language,
real Heading 1-4 styles for navigation, and no images, tables or text boxes.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from ..model import HEADINGS
from . import marker_text, output_blocks, wants_markers

LANG = "ar-EG"

PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
]
RPR_ORDER = [
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
    "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
    "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing",
    "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
    "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
]
SECT_ORDER = [
    "w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr", "w:type",
    "w:pgSz", "w:pgMar", "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType",
    "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
    "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange",
]


def put(parent, tag, order, **attrs):
    """Replace-or-insert child `tag` of `parent` at its schema position."""
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    el = OxmlElement(tag)
    for key, val in attrs.items():
        el.set(qn("w:" + key), str(val))
    later = {qn(t) for t in order[order.index(tag) + 1:]}
    for child in parent:
        if child.tag in later:
            child.addprevious(el)
            break
    else:
        parent.append(el)
    return el


class _Styler:
    def __init__(self, opts):
        self.font = opts.font
        self.base = opts.base_size
        self.line = opts.line_spacing

    def rpr(self, rpr, size=None, bold=None, color=None):
        put(rpr, "w:rFonts", RPR_ORDER, ascii=self.font, hAnsi=self.font, cs=self.font,
            eastAsia=self.font)
        if bold is not None:
            put(rpr, "w:b", RPR_ORDER, val="1" if bold else "0")
            put(rpr, "w:bCs", RPR_ORDER, val="1" if bold else "0")
        if color:
            put(rpr, "w:color", RPR_ORDER, val=color)
        if size:
            put(rpr, "w:sz", RPR_ORDER, val=int(size * 2))
            put(rpr, "w:szCs", RPR_ORDER, val=int(size * 2))
        put(rpr, "w:rtl", RPR_ORDER)
        put(rpr, "w:lang", RPR_ORDER, val=LANG, bidi=LANG)

    def ppr(self, ppr, align=None, before=0, after=2, line=None, keep_next=False):
        if keep_next:
            put(ppr, "w:keepNext", PPR_ORDER)
        put(ppr, "w:bidi", PPR_ORDER)
        spacing = {"before": int(before * 20), "after": int(after * 20)}
        if line:
            spacing.update(line=int(line * 240), lineRule="auto")
        put(ppr, "w:spacing", PPR_ORDER, **spacing)
        if align:
            put(ppr, "w:jc", PPR_ORDER, val=align)


def _make_style(doc, name, base="Normal"):
    styles = doc.styles
    if name in [s.name for s in styles]:
        return styles[name]
    st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = styles[base]
    st.quick_style = True
    return st


def _setup(doc, opts):
    s = _Styler(opts)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    put(sec._sectPr, "w:bidi", SECT_ORDER)

    defaults = doc.styles.element.find(qn("w:docDefaults"))
    s.rpr(defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")))

    normal = doc.styles["Normal"]
    s.rpr(normal.element.get_or_add_rPr(), size=s.base, color="000000")
    s.ppr(normal.element.get_or_add_pPr(), after=2, line=s.line)

    heading_specs = {
        "Heading 1": (s.base + 7, "1F3864", "center", 4, 4),
        "Heading 2": (s.base + 4, "1F3864", "center", 4, 4),
        "Heading 3": (s.base + 2, "833C0B", None, 8, 3),
        "Heading 4": (s.base + 0.5, "1F3864", None, 4, 2),
    }
    for name, (size, color, align, before, after) in heading_specs.items():
        st = doc.styles[name]
        rpr = st.element.get_or_add_rPr()
        for t in ("w:i", "w:iCs"):
            for old in rpr.findall(qn(t)):
                rpr.remove(old)
        s.rpr(rpr, size=size, bold=True, color=color)
        s.ppr(st.element.get_or_add_pPr(), align=align, before=before, after=after,
              keep_next=True)

    box = _make_style(doc, "Box Note")
    s.rpr(box.element.get_or_add_rPr(), size=s.base)
    ppr = box.element.get_or_add_pPr()
    bdr = put(ppr, "w:pBdr", PPR_ORDER)
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + side)
        for k, v in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", "808080")):
            el.set(qn("w:" + k), v)
        bdr.append(el)
    put(ppr, "w:shd", PPR_ORDER, val="clear", color="auto", fill="F2F2F2")
    s.ppr(ppr, after=0, line=s.line)
    put(ppr, "w:ind", PPR_ORDER, left=170, right=170)

    gap = _make_style(doc, "Box Gap")
    s.rpr(gap.element.get_or_add_rPr(), size=4)
    s.ppr(gap.element.get_or_add_pPr(), after=0)

    cap = _make_style(doc, "Picture Caption")
    s.rpr(cap.element.get_or_add_rPr(), size=s.base - 1.5, color="595959")
    s.ppr(cap.element.get_or_add_pPr(), after=2, line=s.line)

    marker = _make_style(doc, "Page Marker")
    s.rpr(marker.element.get_or_add_rPr(), size=s.base - 1.5, bold=True, color="7F7F7F")
    s.ppr(marker.element.get_or_add_pPr(), align="center", after=4)


LABEL = re.compile(r"^([^:]{1,45}:)(.*)$")


def _add_par(doc, text, style, label_bold=True, page_break=False):
    p = doc.add_paragraph(style=style)
    if page_break:
        put(p._p.get_or_add_pPr(), "w:pageBreakBefore", PPR_ORDER)
    m = LABEL.match(text) if label_bold else None
    parts = [(text, False)]
    if m and len(m.group(1).split()) <= 6:
        parts = [(m.group(1), True), (m.group(2), False)]
    for chunk, bold in parts:
        if not chunk:
            continue
        run = p.add_run(chunk)
        rpr = run._r.get_or_add_rPr()
        if bold:
            put(rpr, "w:b", RPR_ORDER)
            put(rpr, "w:bCs", RPR_ORDER)
        put(rpr, "w:rtl", RPR_ORDER)
    return p


def write_docx(document, path, opts):
    doc = WordDocument()
    _setup(doc, opts)
    props = doc.core_properties
    props.title = Path(document.source).stem + " (نسخة نصية)"
    props.language = LANG
    props.subject = "نسخة نصية فقط بدون صور لقراءتها على جهاز برايل"

    markers = wants_markers(document, opts)
    first = True
    for page in document.pages:
        if markers:
            _add_par(doc, marker_text(page.number, opts), "Page Marker", label_bold=False,
                     page_break=not first)
        elif not first:
            pass  # Word input without page structure: just continue
        first = False
        prev_box = None
        for block, text in output_blocks(page, opts):
            box = block.box_id if block.kind == "box" else None
            if box is not None and prev_box is not None and box != prev_box:
                doc.add_paragraph(style="Box Gap")
            prev_box = box
            if block.kind in HEADINGS:
                _add_par(doc, text, f"Heading {HEADINGS[block.kind]}", label_bold=False)
            elif block.kind == "box":
                _add_par(doc, text, "Box Note")
            elif block.kind in ("caption", "picture", "notice"):
                _add_par(doc, text, "Picture Caption", label_bold=False)
            else:
                _add_par(doc, text, "Normal")
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
