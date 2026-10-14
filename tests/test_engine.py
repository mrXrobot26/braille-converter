"""Engine tests: line building, text normalisation, markup, writers, and an
accuracy regression on the local reference PDFs (skipped when absent)."""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from nassbraille.config import OutputOptions  # noqa: E402
from nassbraille.model import Block, Document, Page, markup_to_blocks, page_to_markup  # noqa: E402
from nassbraille.pdf.lines import Token, tokens_text  # noqa: E402
from nassbraille.textnorm import clean, normalize  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures" / "saifi"


def visual(*cells, from_font=False, size=10.0, width=5.0):
    """Tokens laid out left to right as they appear on the page."""
    out, x = [], 0.0
    for i, cell in enumerate(cells):
        space = cell == " "
        mark = not space and all(0x064B <= ord(c) <= 0x0652 for c in cell)
        w = 0.0 if mark else width
        out.append(Token(cell, x, x + w, 100.0, size, 0.25, from_font, space, mark, i,
                         x - width / 2 if mark else x + w / 2))
        x += w
    return out


@pytest.mark.parametrize("cells, expected", [
    (("م", "١", "٩", "١", "٩", " ", "م", "ا", "ع"), "عام ١٩١٩م"),
    (("م", "1", "9", "1", "9", " ", "م", "ا", "ع"), "عام 1919م"),
    (("5", "0", "%", " ", "ة", "ب", "س", "ن"), "نسبة 50%"),
    ((")", "1", "8", "8", "2", "-", "1", "9", "5", "2", "(", " ", "ة", "ر", "ت", "ف", "ل", "ا"),
     "الفترة (1882-1952)"),
    (("و", "ه", " ", "W", "i", "n", "d", "o", "w", "s", " ", "1", "0"), "Windows 10 هو"),
    ((":", "١", "س"), "س١:"),
    (("ل", "لا", "خ"), "خلال"),
    (("١", "٩", "٥", "٢", " ", "و", "ي", "ل", "و", "ي", " ", "٢", "٣"), "٢٣ يوليو ١٩٥٢"),
])
def test_rtl_lines(cells, expected):
    assert tokens_text(visual(*cells)).strip() == expected


def test_mirroring_only_for_glyph_identities():
    # the glyphs drawn are the visual shapes: "(" on the left, ")" on the right
    cells = ("(", "1", "9", "1", "9", ")", " ", "م", "ا", "ع")
    assert tokens_text(visual(*cells, from_font=True)).strip() == "عام (1919)"


def test_mark_follows_its_letter():
    toks = visual("ب", "ت", "ك")
    # a damma drawn over the right-most letter (ك)
    toks.append(Token("ُ", 10.0, 10.0, 99.0, 10.0, 0.25, False, False, True, 9, 12.5))
    assert tokens_text(toks).strip() == "كُتب"


def test_missing_space_from_gap():
    toks = visual("ن", "م")
    far = visual("ي", "ف")
    for t in far:                     # a word 20 pt to the left, no space glyph
        t.x0 -= 20
        t.x1 -= 20
    assert tokens_text(far + toks).strip() == "من في"


@pytest.mark.parametrize("raw, expected", [
    ("كلمة ، و (نص )  ثم:", "كلمة، و (نص) ثم:"),
    ("الســيادة", "السيادة"),
    ("حتى آسيا٠", "حتى آسيا."),
    (" بند", "- بند"),
    ("أحمد", "أحمد"),
    ("عُزِّزت", "عُزِّزت"),
])
def test_clean(raw, expected):
    assert clean(raw) == expected


def test_presentation_forms_become_letters():
    assert normalize("ﻻﺍ") == "لاا"


def test_markup_round_trip():
    page = Page(1, [Block("h3", "عنوان"), Block("para", "فقرة"),
                    Block("box", "سطر ١", box_id=0), Block("box", "سطر ٢", box_id=0),
                    Block("box", "برواز تاني", box_id=1), Block("caption", "(صورة: خريطة)"),
                    Block("picture", "(صورة)")])
    back = markup_to_blocks(page_to_markup(page))
    assert [(b.kind, b.text) for b in back] == [(b.kind, b.text) for b in page.blocks]
    assert [b.box_id for b in back if b.kind == "box"] == [0, 0, 1]


def test_docx_writer_round_trip(tmp_path):
    from docx import Document as WordDocument
    from nassbraille.writers.docx_writer import write_docx
    from nassbraille.writers.txt_writer import write_txt
    doc = Document("x.pdf", "pdf", [Page(1, [Block("h2", "الدرس"), Block("para", "نص ـ عادي")]),
                                    Page(2, [Block("box", "برواز", box_id=0),
                                             Block("picture", "(صورة)")])])
    opts = OutputOptions()
    out = write_docx(doc, tmp_path / "a.docx", opts)
    texts = [p.text for p in WordDocument(str(out)).paragraphs if p.text.strip()]
    assert texts == ["صفحة ١", "الدرس", "نص عادي", "صفحة ٢", "برواز"]
    write_txt(doc, tmp_path / "a.txt", opts)
    data = (tmp_path / "a.txt").read_bytes()
    assert data.startswith(b"\xef\xbb\xbf") and b"\r\n" in data
    opts.digits = "western"
    opts.picture_markers = True
    write_txt(doc, tmp_path / "b.txt", opts)
    text = (tmp_path / "b.txt").read_text(encoding="utf-8-sig")
    assert "صفحة 2" in text and "(صورة)" in text


def _scores(pdf, use_tags=True, pages=None):
    from nassbraille.eval import gold_page, output_words, page_scores
    from nassbraille.pdf.reader import read_pdf
    document = read_pdf(pdf, use_tags=use_tags)
    hit = gold = out = 0
    for page in document.pages:
        if pages and page.number not in pages:
            continue
        s = page_scores(gold_page(FIXTURES / "gold" / f"p{page.number:02d}.txt"),
                        output_words(page))
        hit, gold, out = hit + s["hit"], gold + s["gold"], out + s["out"]
    return hit / gold, hit / max(out, 1)


needs_fixtures = pytest.mark.skipif(not (FIXTURES / "original.pdf").exists(),
                                    reason="local reference PDFs not present")


@needs_fixtures
def test_accuracy_normal_pdf():
    recall, precision = _scores(FIXTURES / "final.pdf")
    assert recall >= 0.999 and precision >= 0.999


@needs_fixtures
def test_accuracy_normal_pdf_untagged():
    recall, precision = _scores(FIXTURES / "final.pdf", use_tags=False)
    assert recall >= 0.995 and precision >= 0.995


@needs_fixtures
def test_accuracy_hard_pdf():
    recall, precision = _scores(FIXTURES / "original.pdf")
    assert recall >= 0.970 and precision >= 0.995
